#!/usr/bin/env python3
"""Rewrite RemoteFlow_Kenya.pptx using the SEE Pakistan pitch-deck rules.

    python3 talks/remoteflow-kenya-pitch/generate_remoteflow_pitch.py
"""

from html import escape
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


INK = RGBColor(0x12, 0x18, 0x16)
MUTED = RGBColor(0x5A, 0x63, 0x5F)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PAPER = RGBColor(0xF4, 0xF7, 0xF5)
MINT = RGBColor(0xE3, 0xF4, 0xEC)
WEAK = RGBColor(0xF8, 0xE8, 0xEA)
GREEN = RGBColor(0x0B, 0x7A, 0x4B)
GREEN_DEEP = RGBColor(0x0A, 0x3D, 0x2C)
ORANGE = RGBColor(0xC4, 0x5A, 0x12)
RED = RGBColor(0xC4, 0x34, 0x2A)
BLACK = RGBColor(0x0E, 0x1A, 0x16)
LIGHT = RGBColor(0xFF, 0xFF, 0xFF)

W = Inches(13.333)
H = Inches(7.5)
FONT = "Calibri"


def _set_shape_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _disable_shadow(shape):
    sp_pr = shape._element.spPr
    effect = sp_pr.find(qn("a:effectLst"))
    if effect is not None:
        sp_pr.remove(effect)


def _font(run):
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    latin = rpr.find(qn("a:latin"))
    if latin is None:
        latin = etree.SubElement(rpr, qn("a:latin"))
    latin.set("typeface", FONT)


def _set_run(paragraph, text, *, size, color, bold=False):
    paragraph.clear()
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    _font(run)
    return run


def add_textbox(slide, left, top, width, height, text, *, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(4 if len(lines) > 1 else 0)
        _set_run(p, line, size=size, color=color, bold=bold)
    return box


def add_rect(slide, left, top, width, height, color, *, radius=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(left), Inches(top), Inches(width), Inches(height))
    _set_shape_fill(shape, color)
    _disable_shadow(shape)
    if radius is not None:
        try:
            shape.adjustments[0] = radius
        except Exception:
            pass
    return shape


def paint(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = PAPER
    add_rect(slide, 0, 0, 0.1, 7.5, GREEN)


def add_kicker(slide, text):
    add_textbox(slide, 0.48, 0.26, 12.2, 0.28, text.upper(), size=11, color=GREEN, bold=True)


def add_title(slide, text, *, top=0.5, size=28):
    add_textbox(slide, 0.48, top, 12.3, 0.7, text, size=size, color=INK, bold=True)


def add_footer(slide, index, total):
    add_textbox(slide, 0.48, 7.14, 8.8, 0.24, "REMOTEFLOW  ·  NAIROBI  ·  PRE-SEED  ·  remoteflow.cc", size=10, color=MUTED)
    add_textbox(slide, 10.4, 7.14, 2.4, 0.24, f"{index:02d}  /  {total:02d}", size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    paint(slide)
    return slide


def card(slide, left, top, width, height, color, title, body, *, title_color=INK, body_size=15):
    add_rect(slide, left, top, width, height, color, radius=0.08)
    add_textbox(slide, left + 0.22, top + 0.16, width - 0.44, 0.32, title, size=12, color=title_color, bold=True)
    add_textbox(slide, left + 0.22, top + 0.5, width - 0.44, height - 0.7, body, size=body_size, color=INK)


def build_pptx(out_path: Path):
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    slides = []

    # 1 — One breath
    s = new_slide(prs)
    add_kicker(s, "Understand · Slide 1")
    add_textbox(s, 0.48, 0.52, 12.3, 0.4, "RemoteFlow", size=18, color=GREEN, bold=True)
    add_textbox(
        s, 0.48, 1.0, 12.3, 1.5,
        "Kenyan freelancers receive USD from clients and withdraw KES to M-Pesa in seconds, at a disclosed rate.",
        size=28, color=INK, bold=True,
    )
    add_textbox(
        s, 0.48, 2.7, 12.3, 0.7,
        "Customer + job + outcome. Not an ecosystem. Not “unified wallet.” The business first, ambition second.",
        size=16, color=MUTED,
    )
    add_rect(s, 0.48, 3.55, 12.36, 3.05, MINT, radius=0.08)
    add_textbox(s, 0.74, 3.75, 11.86, 0.3, "WHAT THE PRODUCT DOES", size=12, color=GREEN, bold=True)
    add_textbox(
        s, 0.74, 4.15, 11.86, 2.15,
        "A client pays in USD, GBP, EUR or AED. RemoteFlow converts at an interbank-linked rate and pushes KES to M-Pesa through Safaricom Daraja B2C — already live.\n\nStartup World Cup 2026  ·  Nairobi  ·  Pre-seed",
        size=16, color=INK,
    )
    slides.append(s)

    # 2 — Problem
    s = new_slide(prs)
    add_kicker(s, "Understand · Problem")
    add_title(s, "PayPal already took the money. Then it takes more.")
    add_textbox(
        s, 0.48, 1.28, 12.3, 0.7,
        "A Kenyan freelancer on Upwork, Fiverr or a direct contract gets paid in USD. The painful behavior is cashing out through PayPal.",
        size=16, color=MUTED,
    )
    card(
        s, 0.48, 2.1, 6.0, 4.35, WEAK,
        "WEAK — industry rant",
        "“Cross-border payments are broken.” Fragmented. Inefficient. The investor still cannot see one person stuck on Friday when rent is due.",
        title_color=RED,
    )
    card(
        s, 6.84, 2.1, 6.0, 4.35, MINT,
        "STRONG — one person, one leak",
        "PayPal converts 3–5% below market. On KES 80,000/month that is KES 2,400–4,000 gone — every month.\n\nIt can also freeze a Kenyan account for up to 180 days. Banks need 2–5 days. Rent does not wait.",
        title_color=GREEN,
    )
    slides.append(s)

    # 3 — Solution + why now
    s = new_slide(prs)
    add_kicker(s, "Understand + why now")
    add_title(s, "USD in. Convert. M-Pesa in under five seconds.")
    card(s, 0.48, 1.4, 3.5, 3.05, LIGHT, "BEFORE", "PayPal or Wise.\nBad FX, a freeze, or a 2–5 day bank wait. Then M-Pesa, eventually.")
    add_rect(s, 4.18, 1.7, 4.96, 2.45, BLACK, radius=0.08)
    add_textbox(s, 4.38, 1.9, 4.56, 0.32, "YOUR PRODUCT", size=12, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(s, 4.38, 2.3, 4.56, 1.55, "RemoteFlow wallet\nReceive → convert → Daraja B2C", size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    card(s, 9.34, 1.4, 3.5, 3.05, MINT, "AFTER", "One workflow. KES on the phone in seconds. Price disclosed before the tap.", title_color=GREEN)
    add_rect(s, 0.48, 4.65, 12.36, 2.0, WHITE, radius=0.08)
    add_textbox(s, 0.7, 4.8, 11.9, 0.28, "WHY NOW?", size=12, color=ORANGE, bold=True)
    add_textbox(
        s, 0.7, 5.15, 11.9, 1.25,
        "Daraja B2C is live on the product. Kenya’s remote USD earnings are already large. CBK has a PSP licence path. PayPal’s Kenyan limits are the tax the workforce is already paying.",
        size=16, color=INK,
    )
    slides.append(s)

    # 4 — Market
    s = new_slide(prs)
    add_kicker(s, "Believe · Market")
    add_title(s, "Replace $7.9T with math you can defend")
    card(
        s, 0.48, 1.4, 4.15, 5.15, WEAK,
        "NOT THIS",
        "$7.9T\n\n“Global cross-border payments by 2030.”\n\nBig number. Low information. An investor cannot see who you sell to this year.",
        title_color=RED, body_size=15,
    )
    add_rect(s, 4.88, 1.4, 7.96, 5.15, MINT, radius=0.08)
    add_textbox(s, 5.16, 1.6, 7.4, 0.3, "BOTTOM-UP BEACHHEAD", size=12, color=GREEN, bold=True)
    add_textbox(
        s, 5.16, 2.05, 7.4, 1.35,
        "5,000 Kenyan freelancers already paid in USD  ×  KES 80,000 / month  =  KES 4.8B volume / year",
        size=20, color=INK, bold=True,
    )
    add_textbox(
        s, 5.16, 3.5, 7.4, 1.7,
        "Net take used in the model: ~0.8% after M-Pesa costs → ~KES 38M revenue if that beachhead is fully served.\n\nFirst slice you can actually acquire: freelancer communities already in beta onboarding — 500 users in year one, not 1 million logos.",
        size=15, color=INK,
    )
    add_textbox(
        s, 5.16, 5.4, 7.4, 0.8,
        "Corridor context, not TAM: KES 210B+ remote earnings into Kenya. East Africa is later. Win Nairobi cash-out first.",
        size=14, color=MUTED,
    )
    slides.append(s)

    # 5 — Traction
    s = new_slide(prs)
    add_kicker(s, "Believe · Proof")
    add_title(s, "Traction is what is live. Policies are not traction.")
    items = [
        ("LIVE", "Safaricom Daraja B2C payouts to M-Pesa."),
        ("LIVE", "USD ingestion via PayPal Business API."),
        ("LIVE", "Production stack: NestJS, PostgreSQL, Redis, Vercel. Double-entry ledger, KYC, FX engine."),
        ("FORMED", "RemoteFlow Limited incorporated. Shareholder agreement executed. Ring-fenced accounts drafted for CBK."),
        ("IN MOTION", "Beta onboarding through freelancer communities. CBK PSP engagement started. Safaricom Spark submitted."),
    ]
    for i, (tag, line) in enumerate(items):
        y = 1.38 + i * 1.02
        add_rect(s, 0.48, y, 12.36, 0.92, WHITE, radius=0.08)
        add_rect(s, 0.48, y, 1.55, 0.92, GREEN if tag == "LIVE" else GREEN_DEEP)
        add_textbox(s, 0.48, y + 0.28, 1.55, 0.4, tag, size=12, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        add_textbox(s, 2.2, y + 0.26, 10.3, 0.5, line, size=16, color=INK)
    slides.append(s)

    # 6 — Model + GTM
    s = new_slide(prs)
    add_kicker(s, "Engine")
    add_title(s, "Who pays, how much, how often — then #1 to #1,000")
    card(s, 0.48, 1.35, 4.0, 2.15, WHITE, "WHO PAYS?", "The freelancer, on every payout. RemoteFlow earns when they get paid.", title_color=GREEN)
    card(s, 4.66, 1.35, 4.0, 2.15, WHITE, "HOW MUCH?", "1.5% FX spread · 0.5% deposit · 0.25% transfer. Disclosed before convert. Not “zero fees.”", title_color=ORANGE)
    card(s, 8.84, 1.35, 4.0, 2.15, WHITE, "HOW OFTEN?", "Every client payment. Recurring with the work, not a one-off download.", title_color=RED)
    add_textbox(s, 0.48, 3.68, 12.3, 0.35, "How customer #1 becomes customer #1,000", size=16, color=INK, bold=True)
    for i, step in enumerate(["Community", "KYC", "First payout", "Referral", "Employer / B2B"]):
        x = 0.48 + i * 2.56
        add_rect(s, x, 4.15, 2.4, 0.78, WHITE, radius=0.08)
        add_textbox(s, x, 4.32, 2.4, 0.45, step, size=13, color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(
        s, 0.48, 5.1, 12.3, 1.5,
        "Repeatable motion: freelancer WhatsApp and community groups → KYC → first M-Pesa payout in seconds → invite the next contractor. Year-one model: 500 users, KES 480M volume, ~KES 3.8M FX revenue. HoneyCoin proved the spread model at scale; RemoteFlow owns the last mile to M-Pesa.",
        size=15, color=MUTED,
    )
    slides.append(s)

    # 7 — Competition
    s = new_slide(prs)
    add_kicker(s, "Why you")
    add_title(s, "The real rival is staying on PayPal")
    add_textbox(s, 0.48, 1.25, 12.3, 0.35, "Do not manufacture a feature table where RemoteFlow wins every column.", size=14, color=MUTED)
    rivals = [
        ("Do nothing / PayPal", "Already adopted. Painful, but zero switching cost. This is the default."),
        ("Wise + bank", "Better FX. Still 2–5 days before M-Pesa. Rent is due Friday."),
        ("HoneyCoin / Payd", "Licensed or scaled corridors. Not built as the freelancer’s M-Pesa cash-out."),
        ("Raenest / Grey", "African freelancer wallets. Weak M-Pesa last mile in Kenya."),
    ]
    for i, (name, note) in enumerate(rivals):
        x = 0.48 + i * 3.2
        add_rect(s, x, 1.75, 3.04, 2.55, WHITE, radius=0.08)
        add_textbox(s, x + 0.18, 1.95, 2.68, 0.85, name, size=16, color=INK, bold=True)
        add_textbox(s, x + 0.18, 2.85, 2.68, 1.2, note, size=14, color=MUTED)
    add_rect(s, 0.48, 4.5, 12.36, 2.1, BLACK, radius=0.08)
    add_textbox(
        s, 0.7, 4.7, 11.9, 0.7,
        "The job is not to prove competitors are bad. It is to prove the advantage compounds.",
        size=16, color=WHITE, bold=True,
    )
    add_textbox(
        s, 0.7, 5.45, 11.9, 0.9,
        "M-Pesa B2C already live + PayPal ingestion already live + community distribution + CBK process started. That stack gets harder to copy as payout volume and licence progress accumulate.",
        size=15, color=WHITE,
    )
    slides.append(s)

    # 8 — Team
    s = new_slide(prs)
    add_kicker(s, "Want more · Team")
    add_title(s, "Team ≠ biographies. Show evidence.")
    add_textbox(s, 0.48, 1.28, 12.3, 0.4, "Gilbert Kinyanjui Kamau  ·  Founder", size=18, color=INK, bold=True)
    evidence = [
        "Built Daraja B2C payouts in production — not a slideware integration.",
        "Built PayPal USD ingestion and the FX / ledger / KYC stack himself. No outsourced build to date.",
        "Prior founder: retail and supply businesses. Knows Kenyan cash collection, not just code.",
        "Lived the problem: Kenyan operator moving money, not a tourist researching “Africa fintech.”",
    ]
    for i, line in enumerate(evidence):
        col, row = i % 2, i // 2
        x = 0.48 + col * 6.4
        y = 1.85 + row * 2.2
        add_rect(s, x, y, 6.16, 2.0, WHITE, radius=0.08)
        add_rect(s, x, y, 0.12, 2.0, GREEN)
        add_textbox(s, x + 0.38, y + 0.55, 5.5, 1.1, line, size=16, color=INK)
    slides.append(s)

    # 9 — Ask
    s = new_slide(prs)
    add_kicker(s, "Want more · Ask")
    add_title(s, "Ask ≠ “strategic investors.” $250K, 18 months, three unlocks.")
    rows = [
        ("AMOUNT", "$250K pre-seed. $0 raised. Founder-funded to date."),
        ("RUNWAY", "18 months if the round closes as planned."),
        ("MILESTONES", "CBK PSP licence work, a liquidity/float pool, and the first 1,000 payout users."),
        ("WHAT CAPITAL CHANGES", "Licence and legal 30% · float 35% · community acquisition 20% · engineering 15%."),
        ("WHAT YOU WANT NEXT", "A cheque, or a warm intro to a Kenya payments / CBK-fluent investor. Spark is submitted. Baobab and Antler are next."),
    ]
    for i, (label, body) in enumerate(rows):
        y = 1.35 + i * 0.95
        add_rect(s, 0.48, y, 12.36, 0.85, WHITE, radius=0.08)
        add_rect(s, 0.48, y, 0.12, 0.85, GREEN)
        add_textbox(s, 0.84, y + 0.08, 12.0, 0.26, label, size=11, color=GREEN, bold=True)
        add_textbox(s, 0.84, y + 0.36, 12.0, 0.4, body, size=15, color=INK)
    slides.append(s)

    # 10 — Close
    s = new_slide(prs)
    add_kicker(s, "Close")
    add_textbox(s, 0.48, 0.7, 12.3, 1.4, "Kenyan freelancers should not lose rent money to PayPal.", size=30, color=INK, bold=True)
    add_textbox(
        s, 0.48, 2.2, 12.3, 1.0,
        "RemoteFlow is live to M-Pesa. The raise is $250K to licence, seed float, and take the community from the first payouts to 1,000 users.",
        size=18, color=MUTED,
    )
    add_rect(s, 0.48, 3.4, 12.36, 3.15, GREEN_DEEP, radius=0.08)
    add_textbox(s, 0.74, 3.65, 11.86, 0.35, "GILBERT KINYANJUI KAMAU  ·  FOUNDER", size=12, color=WHITE, bold=True)
    add_textbox(s, 0.74, 4.1, 11.86, 2.1, "gilbert@remoteflow.cc\nremoteflow.cc  ·  Nairobi\n\nIf a slide did not help you understand, believe, or want in — it is gone.", size=18, color=WHITE)
    slides.append(s)

    total = len(slides)
    for i, slide in enumerate(slides, start=1):
        add_footer(slide, i, total)

    prs.save(out_path)
    return out_path


def build_html(out_path: Path):
    html = """<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <title>RemoteFlow Kenya · rewritten pitch</title>
  <link rel="preconnect" href="https://fonts.googleapis.com" />
  <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin />
  <link href="https://fonts.googleapis.com/css2?family=Instrument+Sans:wght@400;600;700&family=Newsreader:opsz,wght@6..72,600&display=swap" rel="stylesheet" />
  <style>
    :root { --ink:#121816; --muted:#5a635f; --green:#0b7a4b; --paper:#f4f7f5; --mint:#e3f4ec; --weak:#f8e8ea; }
    * { box-sizing: border-box; }
    body { margin:0; font-family:"Instrument Sans",system-ui,sans-serif; background:var(--paper); color:var(--ink); }
    .wrap { max-width:980px; margin:0 auto; padding:40px 24px 80px; }
    h1,h2 { font-family:Newsreader,Georgia,serif; }
    .slide { background:#fff; border:1px solid #dbe7e0; border-radius:18px; padding:28px; margin:18px 0; }
    .kicker { font-size:12px; letter-spacing:.12em; text-transform:uppercase; color:var(--green); font-weight:700; }
    .grid { display:grid; gap:12px; }
    .two { grid-template-columns:1fr 1fr; }
    .three { grid-template-columns:1fr 1fr 1fr; }
    .four { grid-template-columns:repeat(4,1fr); }
    .card { background:#fafafa; border-radius:12px; padding:16px; }
    .weak { background:var(--weak); }
    .strong { background:var(--mint); }
    .black { background:#0e1a16; color:#fff; border-radius:12px; padding:20px; }
    @media (max-width:800px) { .two,.three,.four { grid-template-columns:1fr; } }
  </style>
</head>
<body>
  <div class="wrap">
    <p class="kicker">Rewritten RemoteFlow_Kenya.pptx · PPTX cannot preview in chat</p>
    <h1>Kenyan freelancers receive USD and withdraw KES to M-Pesa in seconds.</h1>
    <p>Same facts as your original deck. Different story: understand, believe, want more. Zero-fee claims, the $7.9T TAM, and the win-every-column table are gone.</p>

    <article class="slide"><div class="kicker">01 One breath</div>
      <h2>Customer + job + outcome</h2>
      <div class="card strong"><p>Kenyan freelancers receive USD from clients and withdraw KES to M-Pesa in seconds, at a disclosed rate.</p></div>
    </article>
    <article class="slide"><div class="kicker">02 Problem</div>
      <h2>One painful cash-out</h2>
      <p>PayPal converts 3–5% below market (KES 2,400–4,000 on KES 80k/month) and can freeze a Kenyan account for 180 days.</p>
    </article>
    <article class="slide"><div class="kicker">03 Solution + why now</div>
      <div class="grid three">
        <div class="card"><strong>Before</strong><p>PayPal/Wise → bad FX, freeze, or bank wait.</p></div>
        <div class="black"><strong>Product</strong><p>Receive → convert → Daraja B2C.</p></div>
        <div class="card strong"><strong>After</strong><p>KES on M-Pesa in seconds. Price disclosed.</p></div>
      </div>
    </article>
    <article class="slide"><div class="kicker">04 Market</div>
      <h2>5,000 × KES 80k/month = KES 4.8B volume</h2>
      <p>~0.8% net take → ~KES 38M revenue on that beachhead. Year one is 500 users, not 1 million.</p>
    </article>
    <article class="slide"><div class="kicker">05 Traction</div>
      <p>Live: Daraja B2C, PayPal ingestion, production stack. In motion: beta community, CBK, Spark.</p>
    </article>
    <article class="slide"><div class="kicker">06 Engine</div>
      <p>Who pays: freelancer on each payout. Price: 1.5% FX + 0.5% deposit + 0.25% transfer. GTM: community → KYC → first payout → referral.</p>
    </article>
    <article class="slide"><div class="kicker">07 Competition</div>
      <p>Do nothing/PayPal, Wise+bank, HoneyCoin/Payd, Raenest/Grey. Advantage: live M-Pesa last mile + ingestion + community + CBK process.</p>
    </article>
    <article class="slide"><div class="kicker">08 Team</div>
      <p>Built Daraja and the stack. Prior Kenyan operator. Lived the cash-out problem.</p>
    </article>
    <article class="slide"><div class="kicker">09 Ask</div>
      <p>$250K · 18 months · licence, float, first 1,000 users. Next: a cheque or a CBK-fluent intro.</p>
    </article>
    <article class="slide"><div class="kicker">10 Close</div>
      <p>gilbert@remoteflow.cc · remoteflow.cc</p>
    </article>
  </div>
</body>
</html>
"""
    out_path.write_text(html, encoding="utf-8")
    return out_path


def main():
    folder = Path(__file__).resolve().parent
    pptx_path = build_pptx(folder / "RemoteFlow_Kenya_rewrite.pptx")
    html_path = build_html(folder / "preview.html")
    print(f"Wrote {pptx_path}")
    print(f"Wrote {html_path}")


if __name__ == "__main__":
    main()
